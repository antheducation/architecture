import re, sys, os, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def inline(par, text):
    # handle **bold**, *italic*, `code`
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for t in tokens:
        if not t: continue
        if t.startswith('**') and t.endswith('**'):
            r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            r = par.add_run(t[1:-1]); r.italic = True
        elif t.startswith('`') and t.endswith('`'):
            r = par.add_run(t[1:-1]); r.font.name = 'Consolas'
        else:
            par.add_run(t)

def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]

def build(md_path, out_path):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        if s.startswith('---') and set(s) <= set('-'):
            p = doc.add_paragraph(); p.add_run('_' * 60).font.color.rgb = RGBColor(0x99,0x99,0x99)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=min(lvl, 4))
            h.text = ''
            inline(h, m.group(2))
            i += 1; continue
        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|?$', lines[i+1].strip()):
            header = split_row(s)
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith('|'):
                rows.append(split_row(lines[j])); j += 1
            ncols = len(header)
            t = doc.add_table(rows=1, cols=ncols)
            t.style = 'Table Grid'
            for k, c in enumerate(header):
                cell = t.rows[0].cells[k]
                cell.text = ''
                inline(cell.paragraphs[0], c)
                for r in cell.paragraphs[0].runs: r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for k in range(ncols):
                    cells[k].text = ''
                    inline(cells[k].paragraphs[0], row[k] if k < len(row) else '')
            doc.add_paragraph()
            i = j; continue
        if s.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            inline(p, s[2:])
            i += 1; continue
        if s == '>':
            i += 1; continue
        m = re.match(r'^[-*]\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Bullet'); inline(p, m.group(1)); i += 1; continue
        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Number'); inline(p, m.group(2)); i += 1; continue
        # paragraph: join wrapped lines
        buf = [s]; j = i + 1
        while j < len(lines):
            nx = lines[j].strip()
            if not nx or nx.startswith(('#','|','>','-','*')) or re.match(r'^\d+\.\s', nx): break
            buf.append(nx); j += 1
        p = doc.add_paragraph(); inline(p, ' '.join(buf))
        i = j
    doc.save(out_path)

src = sys.argv[1]; dst = sys.argv[2]
os.makedirs(dst, exist_ok=True)
for f in sorted(glob.glob(os.path.join(src, '*.md'))):
    name = os.path.splitext(os.path.basename(f))[0] + '.docx'
    build(f, os.path.join(dst, name))
    print('OK', name)
